import sys
import json
import openpyxl

xlsx_path = sys.argv[1]
wb = openpyxl.load_workbook(xlsx_path, data_only=True)
sheets = {}
for name in wb.sheetnames:
    ws = wb[name]
    rows = []
    for r in ws.iter_rows(values_only=True):
        if any(c is not None and str(c).strip() != '' for c in r):
            rows.append([str(c) if c is not None else '' for c in r])
    sheets[name] = rows

print(json.dumps(sheets))
