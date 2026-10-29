import os
import openpyxl

os.makedirs(r"d:\Antigravity\AG7\tests\fixtures", exist_ok=True)
fix_dir = r"d:\Antigravity\AG7\tests\fixtures"

# 1. layout_ponto_virgula.csv (Semicolon, Brazilian numbers & dates)
csv1 = """Data da Transação;Descrição do Lançamento;Valor Documento;Categoria Contábil;Nº Documento
05/03/2026;Recebimento de Cliente - Venda 102;1.540,50;Receita Bruta;NF-102
10/03/2026;Pagamento Fornecedor Alpha;-850,20;Custos Mercadorias;DOC-991
15/03/2026;Tarifa Bancária Manutenção;-45,00;Despesas Bancárias;BOL-12
20/03/2026;Rendimento Aplicação Financeira;120,80;Receitas Financeiras;EXT-88
"""
with open(os.path.join(fix_dir, "layout_ponto_virgula.csv"), "w", encoding="utf-8") as f:
    f.write(csv1)

# 2. layout_virgula.csv (Comma, quoted fields with commas, R$ symbol)
csv2 = """Data,Historico,Valor,Categoria
02/03/2026,"Serviços de Consultoria, Parcela 1","R$ 5.000,00",Receita Serviços
12/03/2026,"Compra de Equipamentos, TI","-R$ 1.250,75",Ativo Imobilizado
25/03/2026,"Aluguel Escritório Central","-R$ 2.400,00",Despesas Administrativas
"""
with open(os.path.join(fix_dir, "layout_virgula.csv"), "w", encoding="utf-8") as f:
    f.write(csv2)

# 3. layout_negativos.csv (Testing different representations of negative numbers)
csv3 = """Data;Descricao;Valor
01/03/2026;Padrao com menos antes;-100,50
02/03/2026;Menos no final;200,75-
03/03/2026;Parenteses contabeis;(350,00)
04/03/2026;Indicador de debito D;420,25 D
05/03/2026;Indicador de credito C;900,00 C
"""
with open(os.path.join(fix_dir, "layout_negativos.csv"), "w", encoding="utf-8") as f:
    f.write(csv3)

# 4. layout_linhas_vazias.csv (Blank lines at start, middle, and end)
csv4 = """

Data;Descricao;Valor;Doc
01/03/2026;Item Normal 1;100,00;1

;;;
02/03/2026;Item Normal 2;-50,00;2


"""
with open(os.path.join(fix_dir, "layout_linhas_vazias.csv"), "w", encoding="utf-8") as f:
    f.write(csv4)

# 5. layout_cabecalhos_repetidos.csv (Duplicate headers and empty column name)
csv5 = """Data;Valor;Valor;Descricao;
01/03/2026;150,00;0,00;Descricao Teste;Informacao Extra
02/03/2026;250,00;10,00;Segunda Linha;Observacao
"""
with open(os.path.join(fix_dir, "layout_cabecalhos_repetidos.csv"), "w", encoding="utf-8") as f:
    f.write(csv5)

# 6. layout_multiplas_abas.xlsx (Multiple sheets using openpyxl)
wb = openpyxl.Workbook()
ws1 = wb.active
ws1.title = "Janeiro"
ws1.append(["Data", "Historico", "Valor", "Centro Custo"])
ws1.append(["15/01/2026", "Venda de Produtos Jan", 3500.50, "Comercial"])
ws1.append(["20/01/2026", "Fornecedor Matéria Prima", -1200.00, "Operacional"])

ws2 = wb.create_sheet(title="Fevereiro")
ws2.append(["Data Vencimento", "Descricao Operacao", "Valor Total (R$)", "Conta"])
ws2.append(["10/02/2026", "Prestação de Serviços Fev", 4800.00, "Receita"])
ws2.append(["18/02/2026", "Folha de Pagamento", -3200.50, "Pessoal"])

ws3 = wb.create_sheet(title="Metas_Anuais")
ws3.append(["Indicador", "Meta", "Responsavel"])
ws3.append(["EBITDA", "18%", "Diretoria"])

wb.save(os.path.join(fix_dir, "layout_multiplas_abas.xlsx"))

# 7. layout_invalido_corrompido.csv (No headers or data)
with open(os.path.join(fix_dir, "layout_invalido_corrompido.csv"), "w", encoding="utf-8") as f:
    f.write("   \n\n   \n")

# 8. layout_invalido.xlsx (Corrupted binary file)
with open(os.path.join(fix_dir, "layout_invalido.xlsx"), "wb") as f:
    f.write(b"CORRUPTED_NOT_A_VALID_ZIP_XLSX_STREAM_1234567890")

print("Sucesso: Todas as fixtures de teste criadas em tests/fixtures/")
